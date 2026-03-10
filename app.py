from __future__ import annotations

import asyncio
import io
import logging
import re
import time as _time
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import uvicorn
from fastapi import Depends, FastAPI, Header, HTTPException, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse

from chatbot.config import settings
from chatbot.graph import agent_graph
from chatbot.models import ChatRequest, ChatResponse
from chatbot.session import session_manager
from chatbot.tools.registry import ALL_TOOLS

logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(name)s  %(message)s")
logger = logging.getLogger(__name__)

_SAP_DATE_RE = re.compile(r"^/Date\((-?\d+)(?:[+-]\d+)?\)/$")

# ── Graph timeout ─────────────────────────────────────────────────────────
_GRAPH_TIMEOUT = 120  # 2 minutes max per chat request

# ── Per-session concurrency guard (prevents double-submit race) ───────────
_session_locks: dict[str, asyncio.Lock] = {}
_session_locks_guard = asyncio.Lock()
_MAX_SESSION_LOCKS = 5000  # prevent unbounded growth


async def _get_session_lock(session_id: str) -> asyncio.Lock:
    """Return a per-session asyncio.Lock, creating one if needed."""
    async with _session_locks_guard:
        if session_id not in _session_locks:
            # Prune old locks if too many accumulate
            if len(_session_locks) >= _MAX_SESSION_LOCKS:
                # Remove locks that aren't currently held
                to_remove = [
                    k for k, v in _session_locks.items() if not v.locked()
                ]
                for k in to_remove[:len(to_remove) // 2]:  # remove half
                    del _session_locks[k]
            _session_locks[session_id] = asyncio.Lock()
        return _session_locks[session_id]


def _format_sap_date_string(value: str) -> str:
    text = str(value).strip()
    match = _SAP_DATE_RE.fullmatch(text)
    if not match:
        return value

    ts = int(match.group(1)) / 1000
    dt = datetime.fromtimestamp(ts, tz=timezone.utc)
    if dt.hour == 0 and dt.minute == 0 and dt.second == 0:
        return dt.strftime("%Y-%m-%d")
    return dt.strftime("%Y-%m-%d %H:%M:%S UTC")


def _normalize_sap_dates(value: Any) -> Any:
    if isinstance(value, dict):
        return {k: _normalize_sap_dates(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_normalize_sap_dates(item) for item in value]
    if isinstance(value, str):
        return _format_sap_date_string(value)
    return value


@asynccontextmanager
async def lifespan(_app: FastAPI):
    await session_manager.connect()
    yield
    from chatbot.mcp_client import close_mcp_client
    from chatbot.aggregate_client import close_agg_client
    await close_mcp_client()
    await close_agg_client()
    await session_manager.disconnect()


app = FastAPI(
    title="SAP Chatbot Agent",
    version="0.2.0",
    description="LangGraph-powered chatbot that queries SAP via OData",
    lifespan=lifespan,
)

_cors_origins = [o.strip() for o in settings.CORS_ORIGINS.split(",") if o.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_methods=["GET", "POST", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization", "X-API-Key"],
    allow_credentials=True,
)

FRONTEND_DIR = Path(__file__).parent / "frontend"
app.mount("/static", StaticFiles(directory=str(FRONTEND_DIR)), name="static")
app.mount("/chatbot-ui", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="chatbot-ui")


# ── API key authentication ────────────────────────────────────────────────

async def verify_api_key(x_api_key: str | None = Header(None)) -> None:
    """Verify API key if one is configured. No-op in dev mode (empty key)."""
    if not settings.API_SECRET_KEY:
        return  # Dev mode — auth disabled
    if x_api_key != settings.API_SECRET_KEY:
        raise HTTPException(
            status_code=401,
            detail="Invalid or missing API key. Provide X-API-Key header.",
        )


# ── Redis-backed rate limiter ─────────────────────────────────────────────

_RATE_LIMIT_RPM = 30          # max requests per minute per tenant
_RATE_LIMIT_WINDOW = 60       # window in seconds
_RATE_KEY_PREFIX = "rate:"


async def _check_rate_limit(tenant_id: str) -> bool:
    """Redis-backed sliding window rate limiter. Returns True if allowed."""
    if session_manager._redis is None:
        return True  # Redis not connected — fail open
    key = f"{_RATE_KEY_PREFIX}{tenant_id}"
    now = _time.time()
    window_start = now - _RATE_LIMIT_WINDOW
    try:
        async with session_manager._redis.pipeline(transaction=True) as pipe:
            await pipe.zremrangebyscore(key, "-inf", window_start)
            await pipe.zcard(key)
            await pipe.zadd(key, {str(now): now})
            await pipe.expire(key, _RATE_LIMIT_WINDOW + 5)
            results = await pipe.execute()
        count = results[1]  # zcard result
        return count < _RATE_LIMIT_RPM
    except Exception as exc:
        logger.warning("Rate limit Redis check failed: %s — failing open", exc)
        return True  # Fail open — don't block on Redis errors


# ── Routes ────────────────────────────────────────────────────────────────

@app.get("/", include_in_schema=False)
async def serve_frontend():
    return FileResponse(str(FRONTEND_DIR / "index.html"))


@app.get("/ui-config")
async def get_ui_config():
    return {
        "dashboard_base_url": settings.DASHBOARD_BASE_URL,
    }


@app.post("/session/start")
async def start_session(
    tenant_id: str | None = None,
    user_id: str | None = None,
    _: None = Depends(verify_api_key),
):

    tid = tenant_id or settings.DEV_TENANT_ID
    uid = user_id or settings.DEV_USER_ID
    sid = await session_manager.create_session(tenant_id=tid, user_id=uid)
    return {"session_id": sid, "tenant_id": tid, "user_id": uid}


@app.post("/chat", response_model=ChatResponse)
async def chat(req: ChatRequest, _: None = Depends(verify_api_key)) -> ChatResponse:

    request_id = uuid.uuid4().hex[:12]
    tenant_id = req.tenant_id or settings.DEV_TENANT_ID
    user_id = req.user_id or settings.DEV_USER_ID

    if not await _check_rate_limit(tenant_id):
        raise HTTPException(
            status_code=429,
            detail="Rate limit exceeded. Please wait before sending more requests.",
        )

    session_id = req.session_id
    if session_id:
        session = await session_manager.get_session(session_id)
        if session is None:
            # Session expired or doesn't exist – create fresh
            session_id = await session_manager.create_session(tenant_id, user_id, session_id)
            session = {"messages": [], "pending_tool": "", "pending_params": {}}
    else:
        session_id = await session_manager.create_session(tenant_id, user_id)
        session = {"messages": [], "pending_tool": "", "pending_params": {}}

    logger.info(
        "[%s] /chat tenant=%s user=%s session=%s msg=%s",
        request_id, tenant_id, user_id, session_id, req.message[:100],
    )

    # Prevent concurrent graph executions on the same session (double-submit guard)
    session_lock = await _get_session_lock(session_id)
    if session_lock.locked():
        raise HTTPException(
            status_code=429,
            detail="A request for this session is already in progress. Please wait.",
        )

    async with session_lock:
        return await _execute_chat(req, request_id, tenant_id, user_id, session_id, session)


async def _execute_chat(
    req: ChatRequest,
    request_id: str,
    tenant_id: str,
    user_id: str,
    session_id: str,
    session: dict[str, Any],
) -> ChatResponse:
    """Inner chat handler — runs under per-session lock."""
    state: dict[str, Any] = {
        "request_id": request_id,
        "tenant_id": tenant_id,
        "user_id": user_id,
        "session_id": session_id,
        "user_message": req.message,
        "messages": session.get("messages", []),
        # Resume fields (in case last turn asked user for missing params)
        "tool_name": session.get("pending_tool", ""),
        "tool_params": session.get("pending_params", {}),
        "last_sap_data": session.get("extra", {}).get("last_sap_data", []),
        "sap_data_history": session.get("extra", {}).get("sap_data_history", []),
    }

    try:
        result = await asyncio.wait_for(
            agent_graph.ainvoke(state),
            timeout=_GRAPH_TIMEOUT,
        )
    except asyncio.TimeoutError:
        logger.error("[%s] Graph execution timed out after %ds", request_id, _GRAPH_TIMEOUT)
        raise HTTPException(
            status_code=504,
            detail="Request timed out. The query took too long — please try a more specific question.",
        )
    except Exception:
        logger.exception("[%s] Graph execution failed", request_id)
        raise HTTPException(
            status_code=500,
            detail=f"An internal error occurred. Reference: {request_id}",
        )

    needs_input = bool(result.get("needs_user_input"))
    validation_errors = result.get("validation_errors", [])

    if needs_input:
        reply = result.get("user_input_prompt", "Could you provide more details?")
        # Save pending state so next turn continues from here
        # Append bot's prompt to history so LLM sees it next turn
        messages_to_save = list(result.get("messages", []))
        messages_to_save.append({"role": "assistant", "content": reply})
        await session_manager.save_session(
            session_id=session_id,
            messages=messages_to_save,
            pending_tool=result.get("tool_name", ""),
            pending_params=result.get("tool_params", {}),
            extra={},  # clear stale export data
        )
    elif validation_errors:
        parts = [f"• {e['param']}: {e['message']}" for e in validation_errors]
        reply = "There were some issues with the query parameters:\n\n" + "\n".join(parts)
        reply += "\n\nPlease correct and try again."
        # Append bot's error reply to history
        messages_to_save = list(result.get("messages", []))
        messages_to_save.append({"role": "assistant", "content": reply})
        await session_manager.save_session(
            session_id=session_id,
            messages=messages_to_save,
            # Preserve pending state so user can correct and retry
            pending_tool=result.get("tool_name", ""),
            pending_params=result.get("tool_params", {}),
            extra={},  # clear stale export data
        )
    else:
        reply = result.get("final_response", "")
        # Store SAP data in session for potential Excel export + chart history
        sap_data = _normalize_sap_dates(result.get("sap_response", {}))
        sap_responses = _normalize_sap_dates(result.get("sap_responses", []))
        sap_data_history = result.get("sap_data_history", [])
        extra: dict[str, Any] = {"sap_data_history": sap_data_history}
        # Multi-tool: merge all response data for export
        if sap_responses and len(sap_responses) > 1:
            combined_data: list[dict] = []
            for resp in sap_responses:
                combined_data.extend(resp.get("data", []))
            if combined_data:
                extra["last_sap_data"] = combined_data
        elif sap_data.get("data"):
            extra["last_sap_data"] = sap_data["data"]
        await session_manager.save_session(
            session_id=session_id,
            messages=result.get("messages", []),
            extra=extra,
        )

    export_available = result.get("export_available", False)
    export_url = f"/export/{session_id}" if export_available or (
        result.get("sap_response", {}).get("count", 0) > 0
    ) else None

    return ChatResponse(
        session_id=session_id,
        reply=reply,
        needs_input=needs_input,
        input_prompt=result.get("user_input_prompt") if needs_input else None,
        chart_data=result.get("chart_data"),
        raw_sap=_normalize_sap_dates(result.get("sap_response")) if not needs_input and not validation_errors else None,
        export_url=export_url,
    )


@app.get("/session/{session_id}")
async def get_session(session_id: str):

    session = await session_manager.get_session(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found or expired")
    return {
        "session_id": session_id,
        "tenant_id": session["tenant_id"],
        "user_id": session["user_id"],
        "message_count": len(session["messages"]),
        "has_pending_tool": bool(session["pending_tool"]),
    }


@app.delete("/session/{session_id}")
async def delete_session(session_id: str):
    await session_manager.delete_session(session_id)
    return {"deleted": session_id}


@app.get("/tools")
async def list_tools():
    """Return the catalogue of available SAP tools."""
    return [
        {
            "name": t.name,
            "description": t.description,
            "parameters": [p.model_dump() for p in t.parameters],
        }
        for t in ALL_TOOLS
    ]


@app.get("/export/{session_id}")
async def export_excel(session_id: str, _: None = Depends(verify_api_key)):
    """Generate an Excel file from the last SAP query stored in the session."""
    session = await session_manager.get_session(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found or expired")

    extra = session.get("extra", {})
    sap_data = _normalize_sap_dates(extra.get("last_sap_data", []))
    if not sap_data:
        raise HTTPException(status_code=404, detail="No SAP data available for export")

    # Flatten nested expand arrays (to_Item, to_DeliveryDocumentItem, etc.)
    # into denormalised rows so each child item gets its own Excel row
    flat_data: list[dict] = []
    for row in sap_data:
        scalar = {}
        child_rows: list[dict] = []
        for k, v in row.items():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                for child in v:
                    child_rows.append(
                        {ck: cv for ck, cv in child.items()
                         if ck != "__metadata" and not isinstance(cv, dict)}
                    )
            elif isinstance(v, (dict, list)):
                continue
            else:
                scalar[k] = v
        if child_rows:
            for child in child_rows:
                flat_data.append({**scalar, **child})
        else:
            flat_data.append(scalar)

    # Collect all unique headers across rows
    all_keys: dict[str, None] = {}
    for r in flat_data:
        for k in r:
            all_keys.setdefault(k, None)
    headers = list(all_keys)

    try:
        import openpyxl
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SAP Data"

        if flat_data:
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = openpyxl.styles.Font(bold=True)

            # Data rows
            for row_idx, row_data in enumerate(flat_data, 2):
                for col_idx, key in enumerate(headers, 1):
                    val = row_data.get(key)
                    # Safety: convert any remaining non-primitive to string
                    if isinstance(val, (dict, list)):
                        val = str(val) if val else ""
                    ws.cell(row=row_idx, column=col_idx, value=val)

            # Auto-width columns
            for col_idx, header in enumerate(headers, 1):
                col_letter = get_column_letter(col_idx)
                max_len = max(
                    len(str(header)),
                    *(len(str(row.get(header, ""))) for row in flat_data[:100]),
                )
                ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=sap_export_{session_id[:8]}.xlsx"},
        )
    except ImportError:
        # Fallback to CSV if openpyxl not installed
        import csv
        buf = io.StringIO()
        if flat_data:
            writer = csv.DictWriter(buf, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for r in flat_data:
                # Convert non-primitives to string for CSV
                safe = {k: (str(v) if isinstance(v, (dict, list)) else v) for k, v in r.items()}
                writer.writerow(safe)
        content = buf.getvalue().encode("utf-8")
        return StreamingResponse(
            io.BytesIO(content),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=sap_export_{session_id[:8]}.csv"},
        )


# ── Node name → user-friendly status label ─────────────────────────────
_NODE_STATUS = {
    "parse_and_route": "Understanding your question...",
    "collect_missing": "Checking parameters...",
    "validate": "Validating query...",
    "execute": "Fetching data...",
    "sql_retry": "Refining query...",
    "final_answer": "Generating answer...",
}


@app.websocket("/ws/chat")
async def ws_chat(ws: WebSocket):
    """WebSocket endpoint for streaming chat responses.

    Protocol (JSON messages):
    → Client sends:  {"message": "...", "session_id": "...", "tenant_id": "", "user_id": ""}
    ← Server sends:
        {"type": "status",  "text": "Fetching data..."}        — progress updates
        {"type": "token",   "text": "partial answer text"}      — streamed LLM tokens
        {"type": "result",  "data": {ChatResponse fields}}      — final result with chart/export
        {"type": "error",   "text": "error message"}            — on failure
    """
    await ws.accept()

    try:
        while True:
            raw = await ws.receive_json()
            message = raw.get("message", "").strip()
            if not message:
                await ws.send_json({"type": "error", "text": "Empty message"})
                continue

            tenant_id = raw.get("tenant_id") or settings.DEV_TENANT_ID
            user_id = raw.get("user_id") or settings.DEV_USER_ID
            session_id = raw.get("session_id") or ""
            request_id = uuid.uuid4().hex[:12]

            # Rate limit
            if not await _check_rate_limit(tenant_id):
                await ws.send_json({"type": "error", "text": "Rate limit exceeded. Please wait."})
                continue

            # Session setup
            if session_id:
                session = await session_manager.get_session(session_id)
                if session is None:
                    session_id = await session_manager.create_session(tenant_id, user_id, session_id)
                    session = {"messages": [], "pending_tool": "", "pending_params": {}}
            else:
                session_id = await session_manager.create_session(tenant_id, user_id)
                session = {"messages": [], "pending_tool": "", "pending_params": {}}

            # Concurrency guard
            session_lock = await _get_session_lock(session_id)
            if session_lock.locked():
                await ws.send_json({"type": "error", "text": "A request is already in progress."})
                continue

            async with session_lock:
                await _ws_execute_chat(ws, request_id, message, tenant_id, user_id, session_id, session)

    except WebSocketDisconnect:
        logger.info("WebSocket client disconnected")
    except Exception as exc:
        logger.exception("WebSocket error: %s", exc)
        try:
            await ws.send_json({"type": "error", "text": f"Internal error: {exc}"})
        except Exception:
            pass


async def _ws_execute_chat(
    ws: WebSocket,
    request_id: str,
    message: str,
    tenant_id: str,
    user_id: str,
    session_id: str,
    session: dict[str, Any],
) -> None:
    """Execute the LangGraph pipeline with real streaming LLM tokens via WebSocket."""
    token_queue: asyncio.Queue = asyncio.Queue()
    ws_send_lock = asyncio.Lock()

    ws_disconnected = False

    async def _ws_send(msg: dict) -> None:
        """Thread-safe WS send — prevents interleaved messages.
        Silently stops sending after a disconnect is detected."""
        nonlocal ws_disconnected
        if ws_disconnected:
            return
        try:
            async with ws_send_lock:
                await ws.send_json(msg)
        except Exception:
            ws_disconnected = True

    state: dict[str, Any] = {
        "request_id": request_id,
        "tenant_id": tenant_id,
        "user_id": user_id,
        "session_id": session_id,
        "user_message": message,
        "messages": session.get("messages", []),
        "tool_name": session.get("pending_tool", ""),
        "tool_params": session.get("pending_params", {}),
        "last_sap_data": session.get("extra", {}).get("last_sap_data", []),
        "sap_data_history": session.get("extra", {}).get("sap_data_history", []),
        "stream_ws": True,
        "_ws_token_queue": token_queue,
    }

    try:
        # Run graph + status updates in a background task
        graph_result_holder: dict[str, Any] = {}
        graph_error_holder: list[Exception] = []

        async def _run_graph():
            try:
                last_node = ""
                result: dict[str, Any] = {}
                async for event in agent_graph.astream(state, stream_mode="updates"):
                    for node_name, node_output in event.items():
                        if node_name != last_node:
                            status_text = _NODE_STATUS.get(node_name, f"Processing ({node_name})...")
                            await _ws_send({"type": "status", "text": status_text})
                            last_node = node_name
                        if isinstance(node_output, dict):
                            result.update(node_output)
                graph_result_holder.update(result)
            except Exception as exc:
                graph_error_holder.append(exc)
            finally:
                # Always unblock the token consumer — the graph may end
                # without ever reaching final_answer (e.g. needs_user_input,
                # validation_errors, or exceptions before the LLM call).
                await token_queue.put(None)

        # Start graph execution in background
        graph_task = asyncio.create_task(_run_graph())

        # Consume tokens from queue in foreground — sends to WS in real time
        first_token = True
        while True:
            try:
                token = await asyncio.wait_for(token_queue.get(), timeout=_GRAPH_TIMEOUT)
            except asyncio.TimeoutError:
                await _ws_send({"type": "error", "text": "Request timed out. Try a more specific question."})
                graph_task.cancel()
                await asyncio.gather(graph_task, return_exceptions=True)
                return

            if token is None:
                break  # stream done

            if first_token:
                first_token = False
            await _ws_send({"type": "token", "text": token})

        # Wait for graph to fully complete
        await graph_task

        if graph_error_holder:
            raise graph_error_holder[0]

        result = graph_result_holder

        # Handle non-streamed paths (needs_input, validation_errors, error fallback)
        needs_input = bool(result.get("needs_user_input"))
        validation_errors = result.get("validation_errors", [])
        final_response = result.get("final_response", "")

        if needs_input:
            reply = result.get("user_input_prompt", "Could you provide more details?")
            messages_to_save = list(result.get("messages", []))
            messages_to_save.append({"role": "assistant", "content": reply})
            await session_manager.save_session(
                session_id=session_id,
                messages=messages_to_save,
                pending_tool=result.get("tool_name", ""),
                pending_params=result.get("tool_params", {}),
                extra={},
            )
            # Stream the prompt text
            if reply and first_token:
                for i in range(0, len(reply), 20):
                    await _ws_send({"type": "token", "text": reply[i:i + 20]})
                    await asyncio.sleep(0.02)
        elif validation_errors:
            parts = [f"• {e['param']}: {e['message']}" for e in validation_errors]
            reply = "There were some issues with the query parameters:\n\n" + "\n".join(parts)
            reply += "\n\nPlease correct and try again."
            messages_to_save = list(result.get("messages", []))
            messages_to_save.append({"role": "assistant", "content": reply})
            await session_manager.save_session(
                session_id=session_id,
                messages=messages_to_save,
                pending_tool=result.get("tool_name", ""),
                pending_params=result.get("tool_params", {}),
                extra={},
            )
            for i in range(0, len(reply), 20):
                await _ws_send({"type": "token", "text": reply[i:i + 20]})
                await asyncio.sleep(0.02)
        else:
            # Normal path — tokens already streamed, just save session
            sap_data = _normalize_sap_dates(result.get("sap_response", {}))
            sap_responses = _normalize_sap_dates(result.get("sap_responses", []))
            sap_data_history = result.get("sap_data_history", [])
            extra: dict[str, Any] = {"sap_data_history": sap_data_history}
            if sap_responses and len(sap_responses) > 1:
                combined_data: list[dict] = []
                for resp in sap_responses:
                    combined_data.extend(resp.get("data", []))
                if combined_data:
                    extra["last_sap_data"] = combined_data
            elif sap_data.get("data"):
                extra["last_sap_data"] = sap_data["data"]
            await session_manager.save_session(
                session_id=session_id,
                messages=result.get("messages", []),
                extra=extra,
            )

            # If no tokens were streamed (e.g. error path returned full text),
            # fall back to chunked sending
            if first_token and final_response:
                for i in range(0, len(final_response), 20):
                    await _ws_send({"type": "token", "text": final_response[i:i + 20]})
                    await asyncio.sleep(0.02)

        # Send final result with metadata
        export_available = result.get("export_available", False)
        export_url = f"/export/{session_id}" if export_available or (
            result.get("sap_response", {}).get("count", 0) > 0
        ) else None

        await _ws_send({
            "type": "result",
            "data": {
                "session_id": session_id,
                "needs_input": needs_input,
                "input_prompt": result.get("user_input_prompt") if needs_input else None,
                "chart_data": result.get("chart_data"),
                "export_url": export_url,
            },
        })

    except asyncio.TimeoutError:
        await _ws_send({"type": "error", "text": "Request timed out. Try a more specific question."})
    except Exception as exc:
        logger.exception("[%s] WS graph execution failed: %s", request_id, exc)
        await _ws_send({"type": "error", "text": f"An error occurred: {exc}"})


@app.get("/health")
async def health():
    """Liveness probe — instant, checks only that the process is alive.

    Use this for Kubernetes livenessProbe.  Never call external services
    here; a slow SAP response must not trigger a pod restart.
    """
    return {"status": "ok"}


@app.get("/ready")
async def readiness():
    """Readiness probe — checks Redis connectivity.

    Use this for Kubernetes readinessProbe.  If Redis is down the app
    cannot serve chat requests, so it should be removed from the load
    balancer until Redis recovers.  SAP is intentionally NOT checked
    here because SAP downtime is transient and should be surfaced to
    users as a message, not as a pod restart.
    """
    from fastapi.responses import JSONResponse

    checks: dict[str, str] = {}
    try:
        if session_manager._redis is not None:
            await session_manager._redis.ping()
            checks["redis"] = "ok"
        else:
            checks["redis"] = "not connected"
    except Exception as exc:
        checks["redis"] = f"error: {type(exc).__name__}"

    overall = "ok" if all(v == "ok" for v in checks.values()) else "degraded"
    status_code = 200 if overall == "ok" else 503
    return JSONResponse({"status": overall, "checks": checks}, status_code=status_code)


if __name__ == "__main__":
    uvicorn.run(
        "app:app",
        host=settings.APP_HOST,
        port=settings.APP_PORT,
        reload=True,
    )
